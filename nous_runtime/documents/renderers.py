"""Deterministic renderers from trusted Document IR to DOCX and PDF."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from nous_runtime.documents.models import DocumentIR


class DocumentRenderError(RuntimeError):
    """Raised when an optional renderer is missing or an export cannot complete."""


def _atomic_target(target: Path, writer) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".nous-document-", suffix=target.suffix, dir=target.parent)
    os.close(descriptor)
    temporary_path = Path(temporary)
    try:
        writer(temporary_path)
        if not temporary_path.is_file() or temporary_path.stat().st_size == 0:
            raise DocumentRenderError("renderer produced an empty file")
        os.replace(temporary_path, target)
    finally:
        temporary_path.unlink(missing_ok=True)


def render_docx(document: DocumentIR, target: Path) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise DocumentRenderError("DOCX export requires python-docx") from exc

    def set_font(run, name="Calibri", size=11, *, bold=None, italic=None, color="000000"):
        run.font.name = name
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic

    def configure_style(style, size, color, before, after, *, bold=True, line=1.0):
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line
        style.paragraph_format.keep_with_next = True

    def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tc_mar.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_table_geometry(table):
        table.autofit = False
        width = 9360
        columns = len(table.columns)
        base, remainder = divmod(width, columns)
        widths = [base + (1 if index < remainder else 0) for index in range(columns)]
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
        if tbl_w.getparent() is None:
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(width))
        tbl_w.set(qn("w:type"), "dxa")
        tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
        if tbl_ind.getparent() is None:
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:w"), "120")
        tbl_ind.set(qn("w:type"), "dxa")
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for value in widths:
            grid_col = OxmlElement("w:gridCol")
            grid_col.set(qn("w:w"), str(value))
            grid.append(grid_col)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                cell.width = Inches(widths[index] / 1440)
                tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
                tc_w.set(qn("w:w"), str(widths[index]))
                tc_w.set(qn("w:type"), "dxa")
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def write(path: Path) -> None:
        doc = Document()
        section = doc.sections[0]
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.492)
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.font.size = Pt(11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6 if document.preset != "narrative_proposal" else 8)
        normal.paragraph_format.line_spacing = 1.10 if document.preset == "standard_business_brief" else (1.25 if document.preset == "compact_reference_guide" else 1.333)
        configure_style(doc.styles["Heading 1"], 16, "2E74B5", 16, 8)
        configure_style(doc.styles["Heading 2"], 13, "2E74B5", 12, 6)
        configure_style(doc.styles["Heading 3"], 12, "1F4D78", 8, 4)
        for style_name in ("List Bullet", "List Number"):
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
            style.paragraph_format.left_indent = Inches(0.5)
            style.paragraph_format.first_line_indent = Inches(-0.25)
            style.paragraph_format.space_after = Pt(8 if document.preset == "standard_business_brief" else 4)
            style.paragraph_format.line_spacing = 1.167 if document.preset == "standard_business_brief" else 1.25
        doc.core_properties.title = document.title
        doc.core_properties.subject = document.subtitle
        doc.core_properties.author = document.author
        doc.core_properties.keywords = f"Nous Document IR; {document.schema_version}; {document.document_id}"
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(header.add_run("Nous Professional Document"), size=9, color="6B7280")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(footer.add_run(document.document_id), size=8, color="808080")
        title = doc.add_paragraph()
        title.paragraph_format.space_after = Pt(4)
        set_font(title.add_run(document.title), size=23, bold=True, color="0B2545")
        if document.subtitle:
            subtitle = doc.add_paragraph()
            subtitle.paragraph_format.space_after = Pt(16)
            set_font(subtitle.add_run(document.subtitle), size=13, color="4B5563")
        metadata = doc.add_paragraph()
        metadata.paragraph_format.space_after = Pt(14)
        set_font(metadata.add_run(f"Author: {document.author}  |  Document ID: {document.document_id}"), size=9, color="6B7280")
        for block in document.blocks:
            if block.kind == "paragraph":
                doc.add_paragraph(block.text)
            elif block.kind == "heading":
                doc.add_paragraph(block.text, style=f"Heading {block.level}")
            elif block.kind in {"bullet_list", "numbered_list"}:
                style = "List Bullet" if block.kind == "bullet_list" else "List Number"
                for item in block.items:
                    doc.add_paragraph(item, style=style)
            elif block.kind == "table":
                table = doc.add_table(rows=len(block.rows), cols=len(block.rows[0]))
                table.style = "Table Grid"
                for row_index, values in enumerate(block.rows):
                    for column_index, value in enumerate(values):
                        cell = table.cell(row_index, column_index)
                        cell.text = value
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(0)
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index else WD_ALIGN_PARAGRAPH.CENTER
                            for run in paragraph.runs:
                                set_font(run, size=9.5, bold=row_index == 0)
                        if row_index == 0:
                            shd = OxmlElement("w:shd")
                            shd.set(qn("w:fill"), "F2F4F7")
                            cell._tc.get_or_add_tcPr().append(shd)
                set_table_geometry(table)
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(4)
            elif block.kind == "code":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.right_indent = Inches(0.25)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run(block.text)
                set_font(run, name="Cascadia Mono", size=9, color="1F2937")
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F4F6F9")
                paragraph._p.get_or_add_pPr().append(shd)
            elif block.kind == "citation":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                value = block.text + (f" ({block.url})" if block.url else "")
                set_font(paragraph.add_run(value), size=9, italic=True, color="4B5563")
            elif block.kind == "page_break":
                doc.add_page_break()
        doc.save(path)

    _atomic_target(target, write)


def _pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = (
        ("NousCJK", Path("C:/Windows/Fonts/msyh.ttc")),
        ("NousCJK", Path("C:/Windows/Fonts/simsun.ttc")),
        ("NousUnicode", Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")),
    )
    for name, path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont(name, str(path)))
                return name
            except Exception:
                continue
    return "Helvetica"


def render_pdf(document: DocumentIR, target: Path) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
    except ImportError as exc:
        raise DocumentRenderError("PDF export requires reportlab") from exc

    def write(path: Path) -> None:
        font = _pdf_font()
        styles = getSampleStyleSheet()
        body = ParagraphStyle("NousBody", parent=styles["BodyText"], fontName=font, fontSize=10.5, leading=14, spaceAfter=6, textColor=colors.HexColor("#1F2937"))
        headings = {
            1: ParagraphStyle("NousH1", parent=body, fontSize=16, leading=19, spaceBefore=16, spaceAfter=8, textColor=colors.HexColor("#2E74B5")),
            2: ParagraphStyle("NousH2", parent=body, fontSize=13, leading=16, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor("#2E74B5")),
            3: ParagraphStyle("NousH3", parent=body, fontSize=12, leading=15, spaceBefore=8, spaceAfter=4, textColor=colors.HexColor("#1F4D78")),
        }
        title_style = ParagraphStyle("NousTitle", parent=body, fontSize=23, leading=27, spaceAfter=4, textColor=colors.HexColor("#0B2545"))
        subtitle_style = ParagraphStyle("NousSubtitle", parent=body, fontSize=13, leading=16, spaceAfter=16, textColor=colors.HexColor("#4B5563"))
        citation_style = ParagraphStyle("NousCitation", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#4B5563"), spaceBefore=4, spaceAfter=4)
        list_style = ParagraphStyle("NousList", parent=body, leftIndent=36, firstLineIndent=-18, spaceAfter=5)
        code_style = ParagraphStyle("NousCode", parent=body, fontName="Courier", fontSize=8.5, leading=11, leftIndent=18, rightIndent=18, backColor=colors.HexColor("#F4F6F9"), borderPadding=8, spaceAfter=8)
        story = [Paragraph(escape(document.title), title_style)]
        if document.subtitle:
            story.append(Paragraph(escape(document.subtitle), subtitle_style))
        story.append(Paragraph(escape(f"Author: {document.author} | Document ID: {document.document_id}"), citation_style))
        story.append(Spacer(1, 8))
        for block in document.blocks:
            if block.kind == "paragraph":
                story.append(Paragraph(escape(block.text).replace("\n", "<br/>"), body))
            elif block.kind == "heading":
                story.append(Paragraph(escape(block.text), headings[block.level]))
            elif block.kind in {"bullet_list", "numbered_list"}:
                for index, item in enumerate(block.items, 1):
                    marker = "•" if block.kind == "bullet_list" else f"{index}."
                    story.append(Paragraph(f"{escape(marker)} {escape(item)}", list_style))
            elif block.kind == "table":
                table_data = [[Paragraph(escape(cell), body) for cell in row] for row in block.rows]
                widths = [(6.5 * inch) / len(block.rows[0])] * len(block.rows[0])
                table = Table(table_data, colWidths=widths, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.extend([table, Spacer(1, 8)])
            elif block.kind == "code":
                story.append(Preformatted(block.text, code_style))
            elif block.kind == "citation":
                value = block.text + (f" ({block.url})" if block.url else "")
                story.append(Paragraph(escape(value), citation_style))
            elif block.kind == "page_break":
                story.append(PageBreak())

        def footer(canvas, doc):
            canvas.saveState()
            canvas.setFont(font, 8)
            canvas.setFillColor(colors.HexColor("#6B7280"))
            canvas.drawCentredString(letter[0] / 2, 0.45 * inch, f"{document.document_id} | Page {doc.page}")
            canvas.restoreState()

        pdf = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=0.75 * inch, title=document.title, author=document.author)
        pdf.build(story, onFirstPage=footer, onLaterPages=footer)

    _atomic_target(target, write)


__all__ = ["DocumentRenderError", "render_docx", "render_pdf"]